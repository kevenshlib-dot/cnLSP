"""生成主题词盲评表：多个方案的前 k 个推荐合并去重、打乱顺序，不显示来源。

用法：
  python make_blind_review.py <pilot.jsonl> <suggest_dir> <out.xlsx> <key.json> \
      --systems csl-ensemble csl-llm csl-ds-flash csl-ds-pro --docs 50 --top 5

suggest_dir 下应有 suggest-<系统>.txt（`annif suggest <系统> docs/*.txt` 的输出）。
key.json 记录每个主题词出自哪些方案，评审期间不要给评审人看。
"""

import argparse
import json
import random
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
HEAD_FILL = PatternFill("solid", fgColor="D9E1F2")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
CHOICES = ["可用", "勉强", "错误"]


def parse_suggest(path: Path, top: int) -> dict[int, list[str]]:
    result, current = {}, None
    for line in path.read_text(encoding="utf-8").splitlines():
        m = re.match(r"Suggestions for .*?(\d+)\.txt", line)
        if m:
            current = int(m.group(1))
            result[current] = []
        elif current is not None and line.startswith("<") and "\t" in line:
            if len(result[current]) < top:
                result[current].append(line.split("\t")[1])
    return result


def font(**kw):
    return Font(name=FONT, **kw)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pilot")
    ap.add_argument("suggest_dir", type=Path)
    ap.add_argument("out")
    ap.add_argument("key")
    ap.add_argument("--systems", nargs="+", required=True)
    ap.add_argument("--docs", type=int, default=50)
    ap.add_argument("--top", type=int, default=5)
    ap.add_argument("--seed", type=int, default=2026)
    args = ap.parse_args()

    docs = [json.loads(line) for line in open(args.pilot, encoding="utf-8")][: args.docs]
    sugg = {s: parse_suggest(args.suggest_dir / f"suggest-{s}.txt", args.top) for s in args.systems}
    rng = random.Random(args.seed)

    key, rows = {}, []
    for i, doc in enumerate(docs):
        terms = {}
        for s in args.systems:
            for t in sugg[s].get(i, []):
                terms.setdefault(t, []).append(s)
        order = list(terms)
        rng.shuffle(order)
        key[str(i)] = {"document_id": doc["document_id"], "terms": {t: terms[t] for t in order}}
        rows.append((i, doc["text"], order))

    wb = Workbook()
    # ---- 说明 ----
    ws = wb.active
    ws.title = "说明"
    lines = [
        ("主题词盲评表", font(bold=True, size=14)),
        (f"共 {len(rows)} 篇文献。每篇文献下列出若干候选主题词，它们来自几种不同的自动标引方案，已合并去重并打乱顺序。", None),
        ("请只根据题名和摘要判断每个主题词是否适合用来标引这篇文献，不必考虑它出自哪个方案。", None),
        ("", None),
        ("需要填写的单元格（黄色）", font(bold=True)),
        ("“评审”表的 D 列“判断”：从下拉列表选择 可用 / 勉强 / 错误。", None),
        ("“评审”表的 E 列“备注”：可选，写下理由或更合适的词。", None),
        ("“补充”表的 C 列：这篇文献还缺哪些重要主题词（用顿号分隔），没有可留空。", None),
        ("", None),
        ("判断标准", font(bold=True)),
        ("可用：编目时会直接采用，准确反映文献主题。", None),
        ("勉强：相关但过宽、过窄或表述不够规范，需要修改后才能用。", None),
        ("错误：与文献主题无关，或是“应用”“研究”这类没有标引价值的泛词。", None),
        ("", None),
        ("填写示例（仅示意格式，不在评审表中）", font(bold=True)),
        ("文献序号 1 ｜ 主题词：柴油机 ｜ 判断：可用 ｜ 备注：", None),
        ("文献序号 1 ｜ 主题词：应用 ｜ 判断：错误 ｜ 备注：泛词", None),
        ("", None),
        ("填完后把文件交回；评分脚本 scripts/score_blind_review.py 会按方案汇总。", None),
    ]
    for r, (text, f) in enumerate(lines, 1):
        c = ws.cell(row=r, column=1, value=text)
        c.font = f or font()
    ws.column_dimensions["A"].width = 110

    # ---- 评审 ----
    ws = wb.create_sheet("评审")
    heads = ["文献序号", "题名与摘要", "主题词", "判断", "备注"]
    for c, h in enumerate(heads, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font, cell.fill, cell.border = font(bold=True), HEAD_FILL, BOX
        cell.alignment = Alignment(horizontal="center", vertical="center")
    dv = DataValidation(type="list", formula1='"' + ",".join(CHOICES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    r = 2
    for i, text, order in rows:
        start = r
        for t in order:
            ws.cell(row=r, column=1, value=i + 1)
            ws.cell(row=r, column=3, value=t)
            for c in (4, 5):
                ws.cell(row=r, column=c).fill = INPUT_FILL
            dv.add(f"D{r}")
            for c in range(1, 6):
                cell = ws.cell(row=r, column=c)
                cell.font, cell.border = font(), BOX
                cell.alignment = Alignment(vertical="top", wrap_text=(c in (2, 5)))
            r += 1
        ws.cell(row=start, column=2, value=text)
        if r - 1 > start:
            ws.merge_cells(start_row=start, start_column=2, end_row=r - 1, end_column=2)
            ws.merge_cells(start_row=start, start_column=1, end_row=r - 1, end_column=1)
        ws.cell(row=start, column=1).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=start, column=2).alignment = Alignment(vertical="top", wrap_text=True)
    last = r - 1
    for col, w in zip("ABCDE", (9, 70, 22, 10, 30)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    # ---- 补充 ----
    ws2 = wb.create_sheet("补充")
    for c, h in enumerate(["文献序号", "题名（摘要前 60 字）", "缺少的重要主题词"], 1):
        cell = ws2.cell(row=1, column=c, value=h)
        cell.font, cell.fill, cell.border = font(bold=True), HEAD_FILL, BOX
    for k, (i, text, _) in enumerate(rows, 2):
        ws2.cell(row=k, column=1, value=i + 1)
        ws2.cell(row=k, column=2, value=text[:60])
        ws2.cell(row=k, column=3).fill = INPUT_FILL
        for c in range(1, 4):
            ws2.cell(row=k, column=c).font = font()
            ws2.cell(row=k, column=c).border = BOX
    for col, w in zip("ABC", (9, 70, 40)):
        ws2.column_dimensions[col].width = w

    # ---- 进度（公式） ----
    ws3 = wb.create_sheet("进度")
    rng_d = f"评审!$D$2:$D${last}"
    stats = [
        ("待判断的主题词总数", f"=COUNTA(评审!$C$2:$C${last})"),
        ("已判断", f"=COUNTA({rng_d})"),
        ("完成比例", "=IF(B1=0,0,B2/B1)"),
        ("可用", f'=COUNTIF({rng_d},"可用")'),
        ("勉强", f'=COUNTIF({rng_d},"勉强")'),
        ("错误", f'=COUNTIF({rng_d},"错误")'),
    ]
    for k, (label, formula) in enumerate(stats, 1):
        ws3.cell(row=k, column=1, value=label).font = font(bold=True)
        cell = ws3.cell(row=k, column=2, value=formula)
        cell.font = font()
        if label == "完成比例":
            cell.number_format = "0.0%"
    ws3.cell(row=8, column=1, value="各方案的人工评定准确率由 scripts/score_blind_review.py 计算，"
                                    "方案对应关系保存在单独的 key 文件中，评审结束前不公开。").font = font(italic=True)
    ws3.column_dimensions["A"].width = 24

    # “进度”表的公式没有缓存值，让 Excel / Numbers 打开时自动计算
    wb.calculation.fullCalcOnLoad = True
    wb.save(args.out)
    Path(args.key).write_text(json.dumps(key, ensure_ascii=False, indent=1), encoding="utf-8")
    n_terms = sum(len(o) for _, _, o in rows)
    print(f"docs={len(rows)} terms={n_terms} rows=2..{last}")


if __name__ == "__main__":
    main()
