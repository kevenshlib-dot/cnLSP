"""汇总盲评结果：按方案计算人工评定的准确率与覆盖率。

用法：python score_blind_review.py <填好的.xlsx> <key.json>

- 严格准确率：推荐中判为“可用”的比例
- 宽松准确率：判为“可用”或“勉强”的比例
- 覆盖率：所有方案合起来判为“可用”的词中，本方案推荐到了多少（不含“补充”表里的新词）
"""

import json
import sys
from collections import defaultdict

from openpyxl import load_workbook


def main(xlsx, key_path):
    key = json.load(open(key_path, encoding="utf-8"))
    ws = load_workbook(xlsx, data_only=True)["评审"]
    judged, doc = {}, None
    for row in ws.iter_rows(min_row=2, values_only=True):
        doc = row[0] if row[0] is not None else doc    # 合并单元格只在首行有值
        if row[2] and row[3]:
            judged[(str(int(doc) - 1), row[2])] = row[3]

    stats = defaultdict(lambda: defaultdict(int))
    usable_all = defaultdict(set)
    for d, info in key.items():
        for term, systems in info["terms"].items():
            j = judged.get((d, term))
            if j is None:
                continue
            if j == "可用":
                usable_all[d].add(term)
            for s in systems:
                stats[s]["n"] += 1
                stats[s][j] += 1
    for d, info in key.items():
        for term, systems in info["terms"].items():
            if term in usable_all[d]:
                for s in systems:
                    stats[s]["hit"] += 1
    total_usable = sum(len(v) for v in usable_all.values())

    print(f"已判断 {len(judged)} 个主题词，其中“可用” {total_usable} 个\n")
    print("方案\t已判断\t严格准确率\t宽松准确率\t覆盖率")
    for s, st in sorted(stats.items()):
        n = st["n"] or 1
        print(f"{s}\t{st['n']}\t{st['可用'] / n:.3f}\t{(st['可用'] + st['勉强']) / n:.3f}\t"
              f"{st['hit'] / (total_usable or 1):.3f}")


if __name__ == "__main__":
    main(*sys.argv[1:3])
